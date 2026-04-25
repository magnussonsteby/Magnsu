"""
Pipedrive REST API v1 client.

Authentication: Settings > Personal preferences > API
Header: x-api-token: <TOKEN>
"""
from io import BytesIO
import httpx


# Ordered list of (api_key, display_label) for all standard deal fields.
# user_id is the owner object in v1; person_name / org_name are direct strings.
_FIELD_MAP = [
    # Core
    ("id",                      "ID"),
    ("title",                   "Title"),
    ("status",                  "Status"),
    ("value",                   "Value"),
    ("currency",                "Currency"),
    ("formatted_value",         "Formatted Value"),
    ("weighted_value",          "Weighted Value"),
    ("probability",             "Probability %"),
    # People / orgs
    ("user_id",                 "Owner"),           # object → extract name
    ("person_name",             "Contact"),
    ("org_name",                "Organization"),
    # Pipeline
    ("pipeline_id",             "Pipeline ID"),
    ("stage_id",                "Stage ID"),
    # Dates
    ("expected_close_date",     "Expected Close"),
    ("add_time",                "Created"),
    ("update_time",             "Updated"),
    ("stage_change_time",       "Stage Changed"),
    ("close_time",              "Closed"),
    ("won_time",                "Won"),
    ("first_won_time",          "First Won"),
    ("lost_time",               "Lost"),
    ("rotten_time",             "Rotten Since"),
    ("lost_reason",             "Lost Reason"),
    # Activities
    ("next_activity_date",      "Next Activity Date"),
    ("next_activity_subject",   "Next Activity Subject"),
    ("next_activity_type",      "Next Activity Type"),
    ("last_activity_date",      "Last Activity Date"),
    ("last_incoming_mail_time", "Last Incoming Email"),
    ("last_outgoing_mail_time", "Last Outgoing Email"),
    # Counts
    ("activities_count",        "Activities"),
    ("done_activities_count",   "Done Activities"),
    ("undone_activities_count", "Open Activities"),
    ("notes_count",             "Notes"),
    ("files_count",             "Files"),
    ("email_messages_count",    "Emails"),
    ("followers_count",         "Followers"),
    ("participants_count",      "Participants"),
    ("products_count",          "Products"),
    # Revenue intelligence (when enabled)
    ("acv",                     "ACV"),
    ("arr",                     "ARR"),
    ("mrr",                     "MRR"),
    # Misc
    ("label",                   "Label"),
    ("visible_to",              "Visibility"),
    ("cc_email",                "Deal Email"),
]

# Fields that are redundant or internal — never included in output
_SKIP_FIELDS = {
    "creator_user_id", "owner_id", "person_id", "org_id",
    "org_hidden", "person_hidden", "active", "deleted",
    "next_activity_id", "last_activity_id",
    "next_activity_time", "next_activity_duration", "next_activity_note",
    "weighted_value_currency", "stage_order_nr",
    "reference_activities_count", "formatted_weighted_value",
}

_KNOWN_API_KEYS = {f[0] for f in _FIELD_MAP} | _SKIP_FIELDS


class PipedriveClient:
    def __init__(self, token: str, company_domain: str = "akselera"):
        self.base_v1 = f"https://{company_domain}.pipedrive.com/api/v1"
        self.base_v2 = f"https://{company_domain}.pipedrive.com/api/v2"
        self.headers = {"x-api-token": token, "Accept": "application/json"}

    async def _get(self, url: str, params: dict | None = None) -> dict:
        async with httpx.AsyncClient(timeout=30) as client:
            resp = await client.get(url, params=params or {}, headers=self.headers)
            resp.raise_for_status()
            body = resp.json()
            if not body.get("success"):
                raise ValueError(body.get("error") or body.get("message") or "API error")
            return body

    async def ping(self) -> dict:
        body = await self._get(f"{self.base_v1}/users/me")
        return body.get("data", {})

    async def raw_deals_sample(self) -> dict:
        """Raw first-page response from v1 and v2 for debugging."""
        results = {}
        for label, url in (("v1", f"{self.base_v1}/deals"), ("v2", f"{self.base_v2}/deals")):
            try:
                async with httpx.AsyncClient(timeout=30) as client:
                    resp = await client.get(
                        url, params={"limit": 3, "status": "open"}, headers=self.headers
                    )
                    body = resp.json()
                    data = body.get("data") or []
                    results[label] = {
                        "status_code": resp.status_code,
                        "success": body.get("success"),
                        "data_len": len(data) if isinstance(data, list) else "n/a",
                        "top_keys": list(data[0].keys())[:15] if data else [],
                        "additional_data": body.get("additional_data"),
                        "error": body.get("error"),
                    }
            except Exception as e:
                results[label] = {"error": str(e)}
        return results

    async def fetch_deal_field_labels(self) -> dict[str, str]:
        """Return {field_key: display_name} for all deal fields (including custom fields)."""
        body = await self._get(f"{self.base_v1}/dealFields")
        return {
            f["key"]: f["name"]
            for f in (body.get("data") or [])
            if f.get("key") and f.get("name")
        }

    async def fetch_owners(self) -> list[dict]:
        body = await self._get(f"{self.base_v1}/users")
        return [
            {"id": u["id"], "name": u["name"]}
            for u in (body.get("data") or [])
            if u.get("active_flag")
        ]

    async def fetch_deals_page_v1(
        self,
        status: str = "open",
        owner_id: int | None = None,
        start: int = 0,
        limit: int = 500,
    ) -> dict:
        params: dict = {"status": status, "start": start, "limit": limit}
        if owner_id:
            params["user_id"] = owner_id
        return await self._get(f"{self.base_v1}/deals", params)

    async def fetch_all_deals(
        self,
        status: str = "open",
        owner_id: int | None = None,
        from_date: str | None = None,
        to_date: str | None = None,
    ) -> list[dict]:
        field_labels = await self.fetch_deal_field_labels()
        rows: list[dict] = []
        start = 0
        while True:
            body = await self.fetch_deals_page_v1(
                status=status, owner_id=owner_id, start=start
            )
            data = body.get("data") or []
            if not data:
                break
            for deal in data:
                if from_date or to_date:
                    ts = (deal.get("update_time") or "")[:10]
                    if from_date and ts and ts < from_date:
                        continue
                    if to_date and ts and ts > to_date:
                        continue
                rows.append(_flatten_deal(deal, field_labels))
            pagination = (body.get("additional_data") or {}).get("pagination", {})
            if not pagination.get("more_items_in_collection"):
                break
            start = pagination.get("next_start", start + 500)
        return rows


def _name_of(obj) -> str:
    if not obj:
        return ""
    if isinstance(obj, dict):
        return obj.get("name") or str(obj.get("value") or obj.get("id") or "")
    return str(obj)


def _val(v):
    """Coerce any field value to a serialisable scalar."""
    if v is None:
        return ""
    if isinstance(v, dict):
        return _name_of(v)
    if isinstance(v, list):
        return ", ".join(str(x) for x in v) if v else ""
    return v


def _flatten_deal(d: dict, field_labels: dict[str, str] | None = None) -> dict:
    result = {}
    for api_key, label in _FIELD_MAP:
        result[label] = _val(d.get(api_key))
    # Append custom / unknown scalar fields using proper label from dealFields if available
    for key, val in d.items():
        if key in _KNOWN_API_KEYS or isinstance(val, (dict, list)) or val is None:
            continue
        label = (field_labels or {}).get(key, key)
        result[label] = val
    return result


def build_html_table(rows: list[dict]) -> str:
    if not rows:
        return "<p>No deals found for the selected criteria.</p>"
    headers = list(rows[0].keys())
    header_row = "".join(f"<th>{h}</th>" for h in headers)
    data_rows = "".join(
        "<tr>" + "".join(f"<td>{r.get(h, '')}</td>" for h in headers) + "</tr>"
        for r in rows
    )
    return f"""
<table border="1" cellpadding="6" cellspacing="0"
       style="border-collapse:collapse;font-family:Arial,sans-serif;font-size:13px">
  <thead style="background:#f0f0f5"><tr>{header_row}</tr></thead>
  <tbody>{data_rows}</tbody>
</table>"""


def build_excel_bytes(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Deals"

    if not rows:
        ws.append(["No deals found"])
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    headers = list(rows[0].keys())

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2F5496")
    center = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, row in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(header, "") or "")

    # Auto-width (cap at 60 chars)
    for col_idx, header in enumerate(headers, 1):
        col_vals = [str(rows[r].get(header, "") or "") for r in range(len(rows))]
        max_len = max(len(header), max((len(v) for v in col_vals), default=0))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 60)

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
